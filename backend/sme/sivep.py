"""SRAG total municipal do MS. Um original nacional por versão, sem imputar zeros."""
import hashlib
import io
import json
import math
import re
import zipfile
from datetime import timedelta
from urllib.parse import quote, urlsplit

import openpyxl
from django.db import transaction
from django.utils import timezone

from .collection import CollectionError, MUNICIPIOS_SP, fetch_public, normalized
from .datasets import SIVEP_URL
from .models import ArquivoColeta, RecursoColeta

API = '/api/v4/projects/cgcovid%2Fdados-publicos/repository/'
ROOT = 'Dados unificados/Unificado Srag'
MAX_FILE = 32 * 1024 * 1024
CACHE_AGE = timedelta(hours=1)
HEADER = ('Município', 'Semana Epidemiológica', 'UF *** Contém BR - CUIDADO ***', 'Codigo', 'casos', 'obitos')


def shared_resource(path, limit):
    """Serializa downloads concorrentes. Erros não renovam o cache nem usam versão vencida."""
    url = 'https://gitlab.com' + path
    with transaction.atomic():
        entry, _ = RecursoColeta.objects.get_or_create(url=url)
        entry = RecursoColeta.objects.select_for_update().get(pk=entry.pk)
        if entry.arquivo_id and entry.verificado_em and timezone.now() - entry.verificado_em < CACHE_AGE:
            artifact = ArquivoColeta.objects.get(pk=entry.arquivo_id)
            return artifact, entry.metadados, entry.verificado_em
        body, metadata = fetch_public('gitlab.com', path, limit=limit)
        artifact, _ = ArquivoColeta.objects.get_or_create(sha256=hashlib.sha256(body).hexdigest(), defaults={'conteudo':body})
        entry.arquivo = artifact
        entry.metadados = metadata
        entry.verificado_em = timezone.now()
        entry.save(update_fields=['arquivo','metadados','verificado_em'])
        return artifact, metadata, entry.verificado_em


class SivepAdapter:
    name, disease, dataset = 'sivep', 'srag', 'sivep_srag'

    def __init__(self, disease='srag'):
        if disease != 'srag':
            raise CollectionError('Esta integração coleta somente SRAG total.')

    def validate(self, source):
        url = urlsplit(source.url)
        match = re.fullmatch(r'ano=(\d{4})', url.query)
        if not match or source.url != f'{SIVEP_URL}?ano={match[1]}' or source.tipo != 'xlsx':
            raise CollectionError('Use o cadastro assistido SIVEP-Gripe, com um ano e formato Excel.')
        year = int(match[1])
        r = source.recorte
        if not 2019 <= year <= timezone.localdate().year:
            raise CollectionError('Ano SIVEP-Gripe fora do período permitido.')
        if r.tipo != 'municipio' or r.uf != 'SP' or r.codigo not in MUNICIPIOS_SP or normalized(r.nome) != normalized(MUNICIPIOS_SP[r.codigo]):
            raise CollectionError('Selecione um município válido de SP.')
        return year

    def fetch(self, source):
        year = self.validate(source)
        tree_path = API + 'tree?per_page=100&path=' + quote(ROOT, safe='')
        catalog, _, _ = shared_resource(tree_path, 1024 * 1024)
        try:
            entries = json.loads(bytes(catalog.conteudo))
            if not isinstance(entries, list) or len(entries) >= 100:
                raise ValueError()
            folders = []
            for item in entries:
                if not isinstance(item, dict):
                    raise ValueError()
                match = re.fullmatch(str(year) + r'(?:_(\d{1,2}))?', item.get('name', ''))
                if item.get('type') == 'tree' and match:
                    week = int(match[1]) if match[1] else 0
                    if not 0 <= week <= 53:
                        raise ValueError()
                    folders.append((week, item['name']))
            if not folders:
                raise CollectionError('O provedor ainda não publicou o ano selecionado.')
            folder = max(folders)[1]
        except (ValueError, TypeError, KeyError, UnicodeError) as exc:
            raise CollectionError('Catálogo SIVEP-Gripe inválido ou alterado.') from exc
        path = API + 'files/' + quote(f'{ROOT}/{folder}/srag_total.xlsx', safe='') + '/raw?ref=main'
        artifact, metadata, checked = shared_resource(path, MAX_FILE)
        self.arquivo_original = artifact
        self.raw_body = bytes(artifact.conteudo)
        self.http_metadata = metadata
        self.provenance_extra = {'arquivo_url':'https://gitlab.com'+path, 'pasta':folder,
            'ano':year, 'aba':'Município e Semana', 'arquivo_verificado_em':checked.isoformat(),
            'ausencias':'Sem registro publicado; não equivale a zero casos.', 'cache_maximo_segundos':3600}
        return self.parse(self.raw_body, source, year)

    def parse(self, body, source, year):
        from .processing import week_start
        if len(body) > MAX_FILE:
            raise CollectionError('Arquivo SIVEP-Gripe excede 32 MiB.')
        workbook = None
        try:
            with zipfile.ZipFile(io.BytesIO(body)) as archive:
                info = archive.infolist()
                if len(info) > 200 or sum(x.file_size for x in info) > 256 * 1024 * 1024:
                    raise CollectionError('Planilha expandida excede os limites de leitura.')
            workbook = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=False, keep_links=False)
            sheet = workbook['Município e Semana']
            sheet.reset_dimensions()
            rows = sheet.iter_rows(values_only=True)
            if tuple(next(rows)) != HEADER:
                raise CollectionError('As colunas da planilha SIVEP-Gripe mudaram.')
            codes = {code[:6]:code for code in MUNICIPIOS_SP}
            selected, seen, totals, weeks = {}, set(), {}, set()
            for index, row in enumerate(rows):
                if index >= 400_000 or len(row) != 6:
                    raise CollectionError('Estrutura ou tamanho inesperado na planilha.')
                name, week, uf, code, cases, deaths = row
                if not isinstance(week, str) or not re.fullmatch(r'\d{1,2}', week):
                    raise CollectionError('Semana inválida na planilha.')
                week = int(week)
                week_start(year, week)
                weeks.add(week)
                if uf != 'SP':
                    continue
                if code not in codes or not isinstance(name, str):
                    raise CollectionError('Código municipal de SP desconhecido.')
                for value in (cases, deaths):
                    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or value < 0 or value != int(value) or value > 9_000_000_000:
                        raise CollectionError('Contagem inválida no SIVEP-Gripe.')
                if (code, week) in seen:
                    raise CollectionError('Município/semana duplicado no SIVEP-Gripe.')
                seen.add((code, week))
                totals[week] = totals.get(week, 0) + int(cases)
                if codes[code] == source.recorte.codigo:
                    selected[week] = int(cases)
            if not weeks:
                raise CollectionError('Arquivo sem semanas publicadas.')
            state_sheet = workbook['UF e Semana']
            state_sheet.reset_dimensions()
            state_rows = state_sheet.iter_rows(values_only=True)
            if tuple(next(state_rows)) != ('Semana Epidemiológica', HEADER[2], 'casos', 'obitos'):
                raise CollectionError('Cabeçalho estadual incompatível.')
            state_totals = {}
            for index, row in enumerate(state_rows):
                if index >= 10_000 or len(row) != 4:
                    raise CollectionError('Tabela estadual inválida.')
                if row[1] == 'SP':
                    week = int(row[0])
                    if week in state_totals:
                        raise CollectionError('Semana estadual duplicada.')
                    state_totals[week] = row[2]
            if totals != state_totals:
                raise CollectionError('Totais municipais e estaduais de SP divergem; coleta rejeitada.')
            return [{'ano':year, 'semana':week, 'codigo':source.recorte.codigo,
                     'casos':selected.get(week)} for week in range(1, max(weeks)+1)]
        except CollectionError:
            raise
        except Exception as exc:
            raise CollectionError('Não foi possível validar a planilha oficial SIVEP-Gripe.') from exc
        finally:
            if workbook:
                workbook.close()

    def normalize(self, rows, source):
        from .processing import week_start
        year = self.validate(source)
        if not isinstance(rows, list) or len(rows) > 53:
            raise CollectionError('Série SIVEP-Gripe inválida.')
        result, seen = [], set()
        for row in rows:
            if not isinstance(row, dict) or row.get('ano') != year or row.get('codigo') != source.recorte.codigo:
                raise CollectionError('Ano ou território incompatível com o lote.')
            week, cases = row.get('semana'), row.get('casos')
            if type(week) is not int or week in seen or (cases is not None and (type(cases) is not int or not 0 <= cases <= 9_000_000_000)):
                raise CollectionError('Semana ou contagem inválida no lote.')
            seen.add(week)
            result.append({'ano':year,'semana':week,'data':week_start(year,week),'casos':cases,'estimados':None})
        return sorted(result, key=lambda row:row['data'])
