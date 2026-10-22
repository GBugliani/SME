"""Casos semanais de síndrome gripal por COVID-19 publicados pelo e-SUS Notifica."""
import io
import json
import re
import zipfile
from urllib.parse import quote, urlsplit

import openpyxl
from django.utils import timezone

from .collection import CollectionError, MUNICIPIOS_SP, normalized
from .datasets import ESUS_URL
from .sivep import API, MAX_FILE, shared_resource

ROOT = 'Dados unificados/Unificado e-SUS'
CITY_HEADER = ('estadoIBGE', 'municipio', 'semEpi', 'casos')
STATE_HEADER = ('estadoIBGE', 'semEpi', 'casos')
CITY_ALIASES = {'florinia': '3516101'}


class EsusCovidAdapter:
    name, disease, dataset = 'esus', 'covid', 'esus_covid'

    def __init__(self, disease='covid'):
        if disease != 'covid':
            raise CollectionError('Esta integração coleta somente SG por COVID-19.')

    def validate(self, source):
        match = re.fullmatch(r'ano=(\d{4})', urlsplit(source.url).query)
        if not match or source.url != f'{ESUS_URL}?ano={match[1]}' or source.tipo != 'xlsx':
            raise CollectionError('Use o cadastro assistido e-SUS Notifica, com um ano e formato Excel.')
        year = int(match[1])
        if not 2023 <= year <= timezone.localdate().year:
            raise CollectionError('Ano e-SUS Notifica fora do período publicado.')
        r = source.recorte
        if r.tipo != 'municipio' or r.uf != 'SP' or r.codigo not in MUNICIPIOS_SP or normalized(r.nome) != normalized(MUNICIPIOS_SP[r.codigo]):
            raise CollectionError('Selecione um município válido de SP.')
        return year

    def fetch(self, source):
        year = self.validate(source)
        catalog, _, _ = shared_resource(API + 'tree?per_page=100&path=' + quote(ROOT, safe=''), 1024 * 1024)
        try:
            entries = json.loads(bytes(catalog.conteudo))
            if not isinstance(entries, list) or len(entries) >= 100:
                raise ValueError()
            folders = []
            for item in entries:
                if not isinstance(item, dict):
                    raise ValueError()
                match = re.fullmatch(str(year) + r'(?:_(\d{1,2}))?', item.get('name', ''))
                if item.get('type') == 'tree' and match:
                    week = int(match[1]) if match[1] else 0
                    if not 0 <= week <= 53:
                        raise ValueError()
                    folders.append((week, item['name']))
            if not folders:
                raise CollectionError('O provedor ainda não publicou o ano selecionado de COVID-19.')
            folder = max(folders)[1]
        except (ValueError, TypeError, KeyError, UnicodeError) as exc:
            raise CollectionError('Catálogo e-SUS Notifica inválido ou alterado.') from exc
        path = API + 'files/' + quote(f'{ROOT}/{folder}/tabela_covid_esus_{year}.xlsx', safe='') + '/raw?ref=main'
        artifact, metadata, checked = shared_resource(path, MAX_FILE)
        self.arquivo_original = artifact
        self.raw_body = bytes(artifact.conteudo)
        self.http_metadata = metadata
        self.provenance_extra = {'arquivo_url':'https://gitlab.com'+path, 'pasta':folder,
            'ano':year, 'aba':'Município e Semana', 'arquivo_verificado_em':checked.isoformat(),
            'ausencias':'Sem registro publicado; não equivale a zero casos.', 'cache_maximo_segundos':3600}
        return self.parse(self.raw_body, source, year)

    def parse(self, body, source, year):
        from .processing import week_start
        if len(body) > MAX_FILE:
            raise CollectionError('Arquivo e-SUS Notifica excede 32 MiB.')
        workbook = None
        try:
            with zipfile.ZipFile(io.BytesIO(body)) as archive:
                info = archive.infolist()
                if len(info) > 200 or sum(item.file_size for item in info) > 256 * 1024 * 1024:
                    raise CollectionError('Planilha expandida excede os limites de leitura.')
            workbook = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True, keep_links=False)
            sheet = workbook['Município e Semana']
            sheet.reset_dimensions()
            rows = sheet.iter_rows(values_only=True)
            if tuple(next(rows)) != CITY_HEADER:
                raise CollectionError('As colunas municipais do e-SUS Notifica mudaram.')
            codes = {normalized(name):code for code, name in MUNICIPIOS_SP.items()}
            if len(codes) != len(MUNICIPIOS_SP):
                raise CollectionError('Nomes municipais ambíguos no catálogo.')
            codes.update(CITY_ALIASES)
            selected, totals, weeks, seen = {}, {}, set(), set()
            for index, row in enumerate(rows):
                if index >= 400_000 or len(row) != 4:
                    raise CollectionError('Estrutura ou tamanho inesperado na planilha e-SUS Notifica.')
                state, name, week, cases = row
                if state != 'SP':
                    continue
                if not isinstance(name, str) or normalized(name) not in codes:
                    raise CollectionError('Município paulista não identificado na planilha e-SUS Notifica.')
                if type(week) is not int or type(cases) is not int or cases < 0 or cases > 9_000_000_000:
                    raise CollectionError('Semana ou contagem inválida no e-SUS Notifica.')
                week_start(year, week)
                code = codes[normalized(name)]
                if (code, week) in seen:
                    raise CollectionError('Município/semana duplicado no e-SUS Notifica.')
                seen.add((code, week))
                weeks.add(week)
                totals[week] = totals.get(week, 0) + cases
                if code == source.recorte.codigo:
                    selected[week] = cases
            if not weeks:
                raise CollectionError('Planilha e-SUS Notifica sem semanas publicadas em SP.')
            state_sheet = workbook['UF e Semana']
            state_sheet.reset_dimensions()
            state_rows = state_sheet.iter_rows(values_only=True)
            if tuple(next(state_rows)) != STATE_HEADER:
                raise CollectionError('As colunas estaduais do e-SUS Notifica mudaram.')
            state_totals = {}
            for index, row in enumerate(state_rows):
                if index >= 10_000 or len(row) != 3:
                    raise CollectionError('Tabela estadual e-SUS Notifica inválida.')
                if row[0] == 'SP':
                    if type(row[1]) is not int or type(row[2]) is not int or row[1] in state_totals:
                        raise CollectionError('Semana estadual e-SUS Notifica inválida.')
                    state_totals[row[1]] = row[2]
            if totals != state_totals:
                raise CollectionError('Totais municipais e estaduais de SP divergem; coleta rejeitada.')
            return [{'ano':year, 'semana':week, 'codigo':source.recorte.codigo,
                     'casos':selected.get(week)} for week in range(1, max(weeks)+1)]
        except CollectionError:
            raise
        except Exception as exc:
            raise CollectionError('Não foi possível validar a planilha oficial e-SUS Notifica.') from exc
        finally:
            if workbook:
                workbook.close()

    def normalize(self, rows, source):
        from .processing import week_start
        year = self.validate(source)
        if not isinstance(rows, list) or len(rows) > 53:
            raise CollectionError('Série e-SUS Notifica inválida.')
        result, seen = [], set()
        for row in rows:
            if not isinstance(row, dict) or row.get('ano') != year or row.get('codigo') != source.recorte.codigo:
                raise CollectionError('Ano ou território incompatível com o lote.')
            week, cases = row.get('semana'), row.get('casos')
            if type(week) is not int or week in seen or (cases is not None and (type(cases) is not int or not 0 <= cases <= 9_000_000_000)):
                raise CollectionError('Semana ou contagem inválida no lote.')
            seen.add(week)
            result.append({'ano':year, 'semana':week, 'data':week_start(year, week), 'casos':cases, 'estimados':None})
        return sorted(result, key=lambda row:row['data'])
